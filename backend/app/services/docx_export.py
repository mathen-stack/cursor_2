from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from app.models.schemas import CoverLetter, GenerationOptions, TailoredResume, TemplateStyle


def _set_narrow_margins(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)


def _add_heading_line(doc: Document, text: str, size: int = 11) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0x1A, 0x3A, 0x2A)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)


def build_resume_docx(resume: TailoredResume, options: GenerationOptions) -> bytes:
    doc = Document()
    _set_narrow_margins(doc)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run(resume.full_name)
    run.bold = True
    run.font.size = Pt(20 if options.template != TemplateStyle.minimal else 18)
    run.font.color.rgb = RGBColor(0x1A, 0x3A, 0x2A)
    name.paragraph_format.space_after = Pt(2)

    if resume.headline:
        hl = doc.add_paragraph()
        hl.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = hl.add_run(resume.headline)
        r.italic = True
        r.font.size = Pt(11)
        hl.paragraph_format.space_after = Pt(2)

    contact_parts = [x for x in [resume.location, resume.email, resume.phone, resume.linkedin_url] if x]
    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = contact.add_run(" · ".join(contact_parts))
    c.font.size = Pt(9)
    contact.paragraph_format.space_after = Pt(8)

    _add_heading_line(doc, "Summary")
    summary = doc.add_paragraph(resume.summary)
    for run in summary.runs:
        run.font.size = Pt(10)

    if resume.skills:
        _add_heading_line(doc, "Skills")
        skills = doc.add_paragraph(", ".join(resume.skills))
        for run in skills.runs:
            run.font.size = Pt(10)

    if resume.experiences:
        _add_heading_line(doc, "Experience")
        for exp in resume.experiences:
            header = doc.add_paragraph()
            left = header.add_run(f"{exp.role} — {exp.company}")
            left.bold = True
            left.font.size = Pt(10)
            dates = f"{exp.start_date} – {exp.end_date}"
            if exp.location:
                dates = f"{exp.location} | {dates}"
            right = header.add_run(f"\n{dates}")
            right.font.size = Pt(9)
            header.paragraph_format.space_after = Pt(2)
            for bullet in exp.bullets:
                bp = doc.add_paragraph(bullet, style="List Bullet")
                for run in bp.runs:
                    run.font.size = Pt(10)

    if resume.education:
        _add_heading_line(doc, "Education")
        for ed in resume.education:
            line = f"{ed.degree}"
            if ed.field:
                line += f" in {ed.field}"
            line += f" — {ed.school}"
            p = doc.add_paragraph()
            r = p.add_run(line)
            r.bold = True
            r.font.size = Pt(10)
            meta = " – ".join(filter(None, [ed.location, f"{ed.start_date} – {ed.end_date}".strip(" –")]))
            if meta:
                m = p.add_run(f"\n{meta}")
                m.font.size = Pt(9)
            if ed.details:
                d = doc.add_paragraph(ed.details)
                for run in d.runs:
                    run.font.size = Pt(10)

    if resume.certifications:
        _add_heading_line(doc, "Certifications")
        for cert in resume.certifications:
            text = cert.name
            if cert.issuer:
                text += f" — {cert.issuer}"
            if cert.date:
                text += f" ({cert.date})"
            p = doc.add_paragraph(text, style="List Bullet")
            for run in p.runs:
                run.font.size = Pt(10)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_cover_letter_docx(letter: CoverLetter, resume: TailoredResume) -> bytes:
    doc = Document()
    _set_narrow_margins(doc)

    header = doc.add_paragraph()
    run = header.add_run(resume.full_name)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x1A, 0x3A, 0x2A)

    contact_parts = [x for x in [resume.email, resume.phone, resume.location, resume.linkedin_url] if x]
    contact = doc.add_paragraph(" · ".join(contact_parts))
    for r in contact.runs:
        r.font.size = Pt(9)
    contact.paragraph_format.space_after = Pt(16)

    greet = doc.add_paragraph(letter.greeting)
    greet.paragraph_format.space_after = Pt(10)

    for para in letter.body_paragraphs:
        p = doc.add_paragraph(para)
        p.paragraph_format.space_after = Pt(10)
        for r in p.runs:
            r.font.size = Pt(11)

    closing = doc.add_paragraph(letter.closing)
    closing.paragraph_format.space_before = Pt(6)
    sig = doc.add_paragraph(letter.signature_name or resume.full_name)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
